# -*- coding: utf-8 -*-
"""
acumular.py
-----------
Fusiona el corte diario de DispatchTrack (reportes/despachos_diarios.xls)
dentro de un maestro acumulado (historico/historico_despachos.xlsx),
usando como llave unica N. de Orden + Fecha ruta (upsert):

  - Misma Orden y misma Fecha ruta  -> se actualiza con el dato mas reciente
                                       (queda el cierre de ese dia).
  - Misma Orden en otra Fecha ruta  -> se agrega una fila nueva; la del dia
                                       anterior se conserva intacta.

Columnas agregadas: "Fecha de corte", "Llave" (Orden|Fecha ruta),
"Último intento" (Sí/No) y "Tiempo en cliente (min)". Guarda el resultado como una Tabla de Excel con nombre fijo (tbl_historico),
en la hoja HISTORICO, conservando TODAS las columnas del export. Si algun dia el
export trae columnas nuevas, el maestro las incorpora sin perder las anteriores.
"""

import os
from datetime import datetime
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SNAPSHOT = os.path.join("reportes", "despachos_diarios.xls")
MASTER_DIR = "historico"
MASTER = os.path.join(MASTER_DIR, "historico_despachos.xlsx")
SHEET = "HISTORICO"
TABLE = "tbl_historico"
KEY = "Orden"
CUT_COL = "Fecha de corte"


def _norm_key(s: pd.Series) -> pd.Series:
    """Orden como texto estable (sin decimales) para que sea una llave confiable."""
    num = pd.to_numeric(s, errors="coerce")
    out = num.astype("Int64").astype("string")
    # para valores no numericos, conserva el texto original
    out = out.where(num.notna(), s.astype("string"))
    return out.fillna("").str.strip()


def _unique_headers(cols):
    """Garantiza encabezados unicos y no vacios (requisito de una Tabla de Excel)."""
    seen, result = {}, []
    for i, c in enumerate(cols):
        name = str(c).strip() or f"Columna_{i+1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        result.append(name)
    return result


COL_LLEGADA = "Tiempo de llegada"     # hora en que llega al punto
COL_CIERRE = "Fecha Llegada"          # hora en que cierra la gestion
COL_TIEMPO = "Tiempo en cliente (min)"


def _segundos_del_dia(s: pd.Series) -> pd.Series:
    """Toma solo la HORA de un valor fecha-hora y la devuelve en segundos."""
    dt = pd.to_datetime(s, errors="coerce")
    return dt.dt.hour * 3600 + dt.dt.minute * 60 + dt.dt.second


def _agregar_tiempo_en_cliente(df: pd.DataFrame) -> pd.DataFrame:
    """
    Minutos que el chofer estuvo en el punto: hora de "Fecha Llegada" menos
    hora de "Tiempo de llegada" (solo se compara la hora, no la fecha).
    Queda vacio si falta alguna de las dos horas o si el resultado es negativo.
    La columna se ubica justo despues de "Fecha Llegada".
    """
    if COL_TIEMPO in df.columns:
        df = df.drop(columns=[COL_TIEMPO])
    if COL_LLEGADA not in df.columns or COL_CIERRE not in df.columns:
        return df

    minutos = (_segundos_del_dia(df[COL_CIERRE]) - _segundos_del_dia(df[COL_LLEGADA])) / 60
    minutos = minutos.where(minutos >= 0).round(1)

    pos = df.columns.get_loc(COL_CIERRE) + 1
    df.insert(pos, COL_TIEMPO, minutos)
    return df


DATE_COL = "Fecha ruta"               # dia real en que la orden salio en ruta
LLAVE_COL = "Llave"                   # Orden|Fecha ruta  (ej. 260047863|2026-10-02)
ULT_COL = "Último intento"            # Sí = fila mas reciente de esa Orden
DERIVADAS = [LLAVE_COL, ULT_COL, COL_TIEMPO]


def _norm_fecha(s: pd.Series, respaldo) -> pd.Series:
    """Fecha como texto AAAA-MM-DD; si falta, usa el valor de respaldo."""
    f = pd.to_datetime(s, errors="coerce").dt.strftime("%Y-%m-%d")
    if isinstance(respaldo, pd.Series):
        respaldo = pd.to_datetime(respaldo, errors="coerce").dt.strftime("%Y-%m-%d")
    return f.fillna(respaldo)


def _preparar(df: pd.DataFrame, respaldo_fecha) -> pd.DataFrame:
    """Normaliza Orden y Fecha ruta, quita columnas derivadas y arma la Llave."""
    df = df.drop(columns=[c for c in DERIVADAS if c in df.columns])
    df[KEY] = _norm_key(df[KEY])
    df = df[df[KEY] != ""].copy()
    df[DATE_COL] = _norm_fecha(
        df[DATE_COL] if DATE_COL in df.columns else pd.Series(index=df.index, dtype="object"),
        respaldo_fecha,
    )
    df[LLAVE_COL] = df[KEY] + "|" + df[DATE_COL].fillna("")
    return df


def _agregar_ultimo_intento(df: pd.DataFrame) -> pd.DataFrame:
    """Marca 'Sí' en la fila de la fecha de ruta mas reciente de cada Orden."""
    ultima = df.groupby(KEY)[DATE_COL].transform("max")
    df[ULT_COL] = (df[DATE_COL] == ultima).map({True: "Sí", False: "No"})
    return df


def _ordenar_columnas(df: pd.DataFrame) -> pd.DataFrame:
    """Fecha de corte, Llave y Último intento al inicio; el resto como venia."""
    primeras = [c for c in [CUT_COL, LLAVE_COL, ULT_COL] if c in df.columns]
    return df[primeras + [c for c in df.columns if c not in primeras]]


def main():
    if not os.path.exists(SNAPSHOT):
        raise SystemExit(f"No se encontro el corte diario: {SNAPSHOT}")

    fecha_corte = datetime.now(ZoneInfo("America/Lima")).strftime("%Y-%m-%d")

    # 1) Corte del dia
    snap = pd.read_excel(SNAPSHOT, sheet_name=0)
    snap.columns = _unique_headers(snap.columns)
    if KEY not in snap.columns:
        raise SystemExit(f"El corte no tiene la columna llave '{KEY}'.")
    snap.insert(0, CUT_COL, fecha_corte)
    snap = _preparar(snap, fecha_corte)
    snap = snap.drop_duplicates(subset=LLAVE_COL, keep="last")

    # 2) Maestro existente (si lo hay)
    if os.path.exists(MASTER):
        master = pd.read_excel(MASTER, sheet_name=SHEET)
        master.columns = _unique_headers(master.columns)
        master = _preparar(master, master[CUT_COL] if CUT_COL in master.columns else fecha_corte)
    else:
        master = pd.DataFrame(columns=snap.columns)

    # 3) Union de columnas: conserva el orden del maestro y agrega nuevas al final
    cols = list(master.columns)
    for c in snap.columns:
        if c not in cols:
            cols.append(c)
    master = master.reindex(columns=cols)
    snap = snap.reindex(columns=cols)

    # 4) Upsert por Orden + Fecha ruta: dentro del mismo dia gana el corte mas
    #    reciente; si la orden sale otro dia, se agrega una fila nueva.
    antes = len(master)
    combinado = pd.concat([master, snap], ignore_index=True)
    combinado = combinado.drop_duplicates(subset=LLAVE_COL, keep="last").reset_index(drop=True)
    combinado = combinado.sort_values([DATE_COL, KEY], na_position="last").reset_index(drop=True)

    combinado = _agregar_ultimo_intento(combinado)
    combinado = _agregar_tiempo_en_cliente(combinado)
    combinado = _ordenar_columnas(combinado)

    nuevas = len(combinado) - antes
    actualizadas = len(snap) - max(nuevas, 0)

    # 5) Escribir maestro como Tabla de Excel
    os.makedirs(MASTER_DIR, exist_ok=True)
    tmp = MASTER + ".tmp.xlsx"
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        combinado.to_excel(xw, index=False, sheet_name=SHEET)

    wb = load_workbook(tmp)
    ws = wb[SHEET]
    nrows, ncols = combinado.shape
    ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
    tbl = Table(displayName=TABLE, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tbl)

    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    for j in range(1, ncols + 1):
        c = ws.cell(1, j)
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(vertical="center")
    for col_cells in ws.iter_cols(min_row=2, max_row=min(nrows + 1, 400)):
        for c in col_cells:
            c.font = Font(name="Arial", size=10)
    ws.freeze_panes = "A2"

    wb.save(MASTER)
    wb.close()
    os.remove(tmp)

    print(f"Fecha de corte : {fecha_corte}")
    print(f"Corte del dia  : {len(snap)} filas (orden + fecha ruta)")
    print(f"Maestro antes  : {antes} filas")
    print(f"  nuevas       : {max(nuevas,0)}")
    print(f"  actualizadas : {max(actualizadas,0)}")
    print(f"Maestro ahora  : {nrows} filas ({combinado[KEY].nunique()} ordenes) x {ncols} columnas")
    print(f"Guardado en    : {MASTER} (hoja {SHEET}, tabla {TABLE})")


if __name__ == "__main__":
    main()
